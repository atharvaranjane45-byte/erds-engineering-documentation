#!/usr/bin/env python3
"""Generate a formatted DOCX from Markdown and a JSON metadata file.

Supported Markdown: headings, paragraphs, simple tables, lists, inline bold/italic/code,
and fenced code blocks. This intentionally avoids executing any content from the input.
"""
from __future__ import annotations
import argparse
import json
import re
from pathlib import Path
from datetime import datetime

try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_SECTION
except ImportError as exc:
    raise SystemExit("Missing dependency. Run: python -m pip install -r requirements.txt") from exc


def add_inline(paragraph, text: str):
    # Small safe inline parser for **bold**, *italic*, and `code`.
    pattern = re.compile(r"(\*\*.*?\*\*|\*.*?\*|`.*?`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("*"):
            paragraph.add_run(token[1:-1]).italic = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_cell_text(cell, value):
    cell.text = ""
    add_inline(cell.paragraphs[0], value.strip())


def set_page_layout(doc):
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.75)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for name, size in [("Title", 24), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)]:
        styles[name].font.name = "Aptos Display"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True


def add_cover(doc, meta):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(70)
    r = p.add_run(str(meta.get("document_type", "ENGINEERING REPORT")).upper())
    r.bold = True
    r.font.size = Pt(13)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(30)
    r = p.add_run(str(meta.get("title", "Untitled Engineering Document")))
    r.bold = True
    r.font.size = Pt(22)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    p.add_run("Submitted by\n").bold = True
    names = meta.get("student_names") or []
    p.add_run(", ".join(names) if names else "[Student name]")
    rolls = meta.get("roll_numbers") or []
    if rolls:
        p.add_run("\nRoll No.: " + ", ".join(map(str, rolls)))
    p.add_run("\n\n" + str(meta.get("department") or "[Department / Branch]"))
    p.add_run("\n" + str(meta.get("college") or "[College / University]"))
    if meta.get("subject"):
        p.add_run("\n\nSubject: " + str(meta["subject"]))
    p.add_run("\n\nProject Guide / Supervisor\n")
    p.add_run(str(meta.get("guide") or "[Guide name, if applicable]"))
    if meta.get("academic_year"):
        p.add_run("\n\nAcademic Year: " + str(meta["academic_year"]))
    if meta.get("date"):
        p.add_run("\n" + str(meta["date"]))
    doc.add_page_break()


def parse_markdown(doc, content):
    lines = content.splitlines()
    i = 0
    in_code = False
    code_lines = []
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if not in_code:
                in_code, code_lines = True, []
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                in_code = False
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        # Markdown table: a header row followed by separator row.
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?[\s:|-]+\|?\s*$", lines[i+1]):
            headers = [x.strip() for x in stripped.strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([x.strip() for x in lines[i].strip().strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = "Light Shading Accent 1"
            for j, h in enumerate(headers):
                add_cell_text(table.rows[0].cells[j], h)
                for run in table.rows[0].cells[j].paragraphs[0].runs:
                    run.bold = True
            for row in rows:
                cells = table.add_row().cells
                for j in range(len(headers)):
                    add_cell_text(cells[j], row[j] if j < len(row) else "")
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            level = min(len(heading.group(1)), 3)
            p = doc.add_heading(level=level)
            add_inline(p, heading.group(2))
            i += 1
            continue

        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        numbered = re.match(r"^\d+[.)]\s+(.*)$", stripped)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, bullet.group(1))
            i += 1
            continue
        if numbered:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, numbered.group(1))
            i += 1
            continue

        # Join wrapped source lines until a blank or block element.
        parts = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or nxt.startswith("#") or nxt.startswith("```") or nxt.startswith("|") or re.match(r"^[-*]\s+", nxt) or re.match(r"^\d+[.)]\s+", nxt):
                break
            parts.append(nxt)
            i += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(parts))


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--markdown", required=True, help="Markdown source report")
    ap.add_argument("--metadata", required=True, help="JSON metadata file")
    ap.add_argument("--output-dir", default="output")
    ap.add_argument("--filename", default="report.docx")
    ap.add_argument("--no-cover", action="store_true", help="Do not generate a cover page")
    args = ap.parse_args()

    md_path = Path(args.markdown)
    meta_path = Path(args.metadata)
    if not md_path.is_file():
        raise SystemExit(f"Markdown input not found: {md_path}")
    if not meta_path.is_file():
        raise SystemExit(f"Metadata JSON not found: {meta_path}")
    try:
        meta = json.loads(meta_path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise SystemExit(f"Invalid metadata JSON: {exc}") from exc

    out_dir = Path(args.output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    doc = Document()
    set_page_layout(doc)
    if not args.no_cover:
        add_cover(doc, meta)
    parse_markdown(doc, md_path.read_text(encoding="utf-8"))
    out = out_dir / args.filename
    doc.save(out)
    print(f"Created DOCX: {out.resolve()}")


if __name__ == "__main__":
    main()
